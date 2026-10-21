"""Banka/kart hesap dökümü okuma ve kurala göre kategori eşleme (saf fonksiyonlar).

Bankalar farklı kolon adları, tarih biçimleri ve tutar düzenleri kullanır. Kullanıcı kolonları
seçer; tutar ya tek bir işaretli kolondan (eksi = gider) ya da ayrı Borç/Alacak kolonlarından okunur.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

from divitrack.core.exceptions import ImportFormatError, ValidationError
from divitrack.core.money import ZERO, to_decimal

DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%y", "%Y/%m/%d", "%d.%m.%Y %H:%M:%S")


def tr_lower(text: str) -> str:
    """Türkçe büyük/küçük harf dönüşümü (I→ı, İ→i) ve boşluk sadeleştirme."""
    return " ".join(text.replace("I", "ı").replace("İ", "i").lower().split())


def parse_date(text: str) -> date:
    value = text.strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    raise ValidationError(f"Tarih anlaşılamadı: {text!r}")


def decode(data: str | bytes) -> str:
    """UTF-8 (BOM'lu/BOM'suz) dener, olmazsa Türkçe Windows kod sayfasına düşer."""
    if isinstance(data, str):
        return data
    for encoding in ("utf-8-sig", "cp1254"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def sniff_delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:20])
    counts = {d: sample.count(d) for d in (";", ",", "\t", "|")}
    return max(counts, key=lambda d: counts[d]) if any(counts.values()) else ","


XLSX_MAGIC = b"PK\x03\x04"
XLS_MAGIC = b"\xd0\xcf\x11\xe0"


def _cell(value: object) -> str:
    """Excel hücresini metne çevirir. Sayılar Türkçe ondalıkla yazılır ki "15.125" binlik sanılmasın."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int | float | Decimal):
        text = format(Decimal(str(value)).normalize(), "f")
        return text.replace(".", ",")
    return " ".join(str(value).split())


def _xlsx_rows(data: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover — requirements.txt'te var
        raise ImportFormatError("Excel dosyaları için openpyxl kurulu olmalı") from exc
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFormatError(f"Excel dosyası okunamadı: {exc}") from exc
    try:
        sheet = wb.worksheets[0]
        return [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        wb.close()


def read_table(data: str | bytes) -> tuple[list[str], list[list[str]]]:
    """CSV veya Excel (.xlsx, ilk sayfa) dökümünü başlık + satırlar olarak okur. Başlık, en az iki
    dolu hücresi olan ilk satırdır (bankaların dökümün başına koyduğu hesap bilgisi satırları atlanır)."""
    if isinstance(data, bytes) and data.startswith(XLS_MAGIC):
        raise ImportFormatError("Eski .xls biçimi desteklenmiyor; dosyayı Excel'de .xlsx veya CSV olarak kaydedin")
    if isinstance(data, bytes) and data.startswith(XLSX_MAGIC):
        rows = _xlsx_rows(data)
    else:
        text = decode(data)
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=sniff_delimiter(text))]
    rows = [r for r in rows if any(r)]
    width = max((len(r) for r in rows), default=0)
    start = next((i for i, r in enumerate(rows) if sum(1 for c in r if c) >= min(2, width) and len(r) == width), None)
    if start is None:
        raise ImportFormatError("Dosyada tablo bulunamadı")
    header = [h or f"Kolon {i + 1}" for i, h in enumerate(rows[start])]
    return header, [r + [""] * (len(header) - len(r)) for r in rows[start + 1 :]]


@dataclass(frozen=True, slots=True)
class StatementRow:
    line: int
    day: date
    amount: Decimal  # işaretli: eksi = gider, artı = gelir
    description: str

    @property
    def is_expense(self) -> bool:
        return self.amount < 0


@dataclass(frozen=True, slots=True)
class ColumnMap:
    date: int
    description: int
    amount: int | None = None  # tek, işaretli tutar kolonu
    debit: int | None = None  # Borç (gider, pozitif yazılır)
    credit: int | None = None  # Alacak (gelir)
    expense_positive: bool = False  # tek kolonda giderler pozitif yazılıyorsa (kredi kartı dökümleri)

    def validate(self) -> None:
        if self.amount is None and self.debit is None and self.credit is None:
            raise ValidationError("Tutar kolonu veya Borç/Alacak kolonlarından en az biri seçilmeli")


def _amount(cell: str) -> Decimal:
    text = cell.replace("TL", "").replace("₺", "").replace("TRY", "").strip()
    if not text:
        return ZERO
    negative = text.startswith("-") or text.endswith("-") or (text.startswith("(") and text.endswith(")"))
    number = text.strip("-()+ ").replace(" ", "")
    # Türk bankaları binlik ayırıcı olarak nokta kullanır: "15.000" on beş bindir, on beş değil.
    if "," not in number and number.count(".") == 1 and len(number.split(".")[1]) == 3:
        number = number.replace(".", "")
    value = to_decimal(number, field="Tutar")
    return -value if negative else value


def parse_statement(rows: Sequence[Sequence[str]], cols: ColumnMap, *, first_line: int = 2) -> list[StatementRow]:
    """Satırları işaretli tutarlara çevirir. Tarihi veya tutarı olmayan satırlar (ara toplam,
    devreden bakiye) atlanır; tarihi olup tutarı okunamayan satırda hata verilir."""
    cols.validate()
    out: list[StatementRow] = []
    for offset, row in enumerate(rows):
        line = first_line + offset
        raw_date = row[cols.date] if cols.date < len(row) else ""
        if not raw_date.strip():
            continue
        try:
            day = parse_date(raw_date)
        except ValidationError:
            continue  # "Toplam", "Devreden bakiye" gibi satırlar
        try:
            if cols.amount is not None:
                amount = _amount(row[cols.amount])
                if cols.expense_positive:
                    amount = -amount
            else:
                debit = _amount(row[cols.debit]) if cols.debit is not None else ZERO
                credit = _amount(row[cols.credit]) if cols.credit is not None else ZERO
                amount = abs(credit) - abs(debit)
        except (ValidationError, IndexError) as exc:
            raise ImportFormatError(f"Satır {line}: {exc}") from exc
        if amount == 0:
            continue
        description = " ".join(row[cols.description].split()) if cols.description < len(row) else ""
        out.append(StatementRow(line, day, amount, description))
    return out


def match_category(description: str, rules: Sequence[tuple[str, int]]) -> int | None:
    """İlk eşleşen kuralın kategori id'si. Kural, açıklamada geçen bir metin parçasıdır
    (büyük/küçük harf ve Türkçe I/İ farkı gözetilmez). Uzun kalıplar önce denenir."""
    text = tr_lower(description)
    for pattern, category_id in sorted(rules, key=lambda r: -len(r[0])):
        if pattern and tr_lower(pattern) in text:
            return category_id
    return None
